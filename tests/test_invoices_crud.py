"""Tests des vues factures/clients, notifications et exports."""

from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.urls import reverse

from accounts.models import Company, UserCompany
from core.models import Currency
from invoices.models import Client, Invoice, InvoiceStatusHistory
from notifications.models import Notification

User = get_user_model()


@pytest.fixture
def company(db):
    return Company.objects.create(
        name="ACME",
        identifier="12345678900011",
        address="1 rue test",
        postal_code="75001",
        city="Paris",
        contact_email="acme@example.com",
    )


@pytest.fixture
def user(company):
    u = User.objects.create_user(
        username="stan", email="stan@example.com", password="pass12345", role="admin"
    )
    UserCompany.objects.create(user=u, company=company, is_primary=True)
    return u


@pytest.fixture
def auth_client(client, user):
    client.force_login(user)
    return client


@pytest.fixture
def eur(db):
    return Currency.objects.create(code="EUR", name="Euro", symbol="€")


@pytest.fixture
def customer(company):
    return Client.objects.create(
        company=company,
        name="Client SA",
        identifier="98765432100022",
        address="2 av. test",
        postal_code="69001",
        city="Lyon",
    )


@pytest.fixture
def invoice(company, customer, eur, user):
    return Invoice.objects.create(
        company=company,
        client=customer,
        invoice_number="F-2026-001",
        platform="chorus_pro",
        amount_ht=Decimal("1000.00"),
        vat_amount=Decimal("200.00"),
        amount_ttc=Decimal("1200.00"),
        currency=eur,
        issue_date=date.today() - timedelta(days=60),
        due_date=date.today() - timedelta(days=10),
        status="submitted",
        created_by=user,
    )


def test_invoice_list_requires_login(client):
    response = client.get(reverse("invoices:invoice_list"))
    assert response.status_code == 302
    assert "/accounts/login/" in response.url


def test_invoice_list_and_filters(auth_client, invoice):
    response = auth_client.get(reverse("invoices:invoice_list"))
    assert response.status_code == 200
    assert b"F-2026-001" in response.content

    response = auth_client.get(reverse("invoices:invoice_list"), {"status": "paid"})
    assert b"F-2026-001" not in response.content

    response = auth_client.get(reverse("invoices:invoice_list"), {"overdue": "1"})
    assert b"F-2026-001" in response.content


def test_create_client_then_invoice(auth_client, company, eur):
    response = auth_client.post(
        reverse("invoices:client_create"),
        {
            "name": "Nouveau client",
            "identifier": "11122233300044",
            "address": "3 bd test",
            "postal_code": "33000",
            "city": "Bordeaux",
            "country": "France",
        },
    )
    assert response.status_code == 302
    customer = Client.objects.get(identifier="11122233300044")
    assert customer.company == company

    response = auth_client.post(
        reverse("invoices:invoice_create"),
        {
            "client": customer.id,
            "invoice_number": "F-2026-002",
            "platform": "manual",
            "amount_ht": "500",
            "amount_ttc": "600",
            "currency": eur.id,
            "issue_date": date.today().isoformat(),
            "due_date": (date.today() + timedelta(days=30)).isoformat(),
        },
    )
    assert response.status_code == 302
    created = Invoice.objects.get(invoice_number="F-2026-002")
    assert created.vat_amount == Decimal("100.00")  # TVA calculée
    assert created.company == company
    assert created.is_compliant is False
    assert created.days_overdue == 0


def test_invoice_form_rejects_duplicate_number_and_bad_dates(
    auth_client, invoice, customer, eur
):
    response = auth_client.post(
        reverse("invoices:invoice_create"),
        {
            "client": customer.id,
            "invoice_number": "F-2026-001",
            "platform": "manual",
            "amount_ht": "100",
            "amount_ttc": "90",
            "currency": eur.id,
            "issue_date": date.today().isoformat(),
            "due_date": (date.today() - timedelta(days=1)).isoformat(),
        },
    )
    assert response.status_code == 200
    form = response.context["form"]
    assert "invoice_number" in form.errors
    assert "amount_ttc" in form.errors
    assert "due_date" in form.errors


def test_invoice_detail_shows_missing_documents_and_overdue(auth_client, invoice):
    response = auth_client.get(reverse("invoices:invoice_detail", args=[invoice.id]))
    assert response.status_code == 200
    content = response.content.decode()
    assert "documents manquants" in content
    assert "Bon de commande" in content
    invoice.refresh_from_db()
    assert invoice.days_overdue == 10


def test_change_status_creates_history_and_payment_date(auth_client, invoice):
    url = reverse("invoices:invoice_change_status", args=[invoice.id])
    response = auth_client.post(url, {"status": "paid", "notes": "Virement reçu"})
    assert response.status_code == 302
    invoice.refresh_from_db()
    assert invoice.status == "paid"
    assert invoice.payment_date == date.today()
    assert invoice.days_overdue == 0
    history = InvoiceStatusHistory.objects.get(invoice=invoice)
    assert history.previous_status == "submitted"
    assert history.new_status == "paid"
    assert history.notes == "Virement reçu"


def test_readonly_user_cannot_create(client, company, eur):
    ro = User.objects.create_user(
        username="ro", email="ro@example.com", password="pass12345", role="readonly"
    )
    UserCompany.objects.create(user=ro, company=company, is_primary=True)
    client.force_login(ro)
    assert client.get(reverse("invoices:invoice_list")).status_code == 200
    assert client.get(reverse("invoices:invoice_create")).status_code == 403


def test_other_company_invoice_is_hidden(auth_client, invoice, eur):
    other = Company.objects.create(
        name="Autre",
        identifier="55566677700088",
        address="x",
        postal_code="00000",
        city="x",
        contact_email="o@example.com",
    )
    other_client = Client.objects.create(
        company=other, name="C2", identifier="1", address="x", postal_code="0", city="x"
    )
    foreign = Invoice.objects.create(
        company=other,
        client=other_client,
        invoice_number="X-1",
        platform="manual",
        amount_ht=1,
        vat_amount=0,
        amount_ttc=1,
        currency=eur,
        issue_date=date.today(),
        due_date=date.today(),
    )
    assert (
        auth_client.get(
            reverse("invoices:invoice_detail", args=[foreign.id])
        ).status_code
        == 404
    )
    response = auth_client.get(reverse("invoices:invoice_list"))
    assert b"X-1" not in response.content


def test_export_csv_and_xlsx(auth_client, invoice):
    url = reverse("reporting:export_invoices")
    response = auth_client.get(url, {"format": "csv"})
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    body = response.content.decode("utf-8-sig")
    assert "F-2026-001" in body
    assert "Client SA" in body

    response = auth_client.get(url, {"format": "xlsx"})
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(response.content)).active
    assert ws["A1"].value == "N° facture"
    assert ws["A2"].value == "F-2026-001"
    assert ws.max_row == 2


def test_invoices_check_command_notifies_once_per_day(invoice, user, auth_client):
    call_command("invoices_check")
    notifications = Notification.objects.filter(user=user)
    codes = set(notifications.values_list("notification_type__code", flat=True))
    assert codes == {"invoice_overdue", "invoice_non_compliant"}
    assert notifications.count() == 2

    call_command("invoices_check")  # pas de doublon le même jour
    assert Notification.objects.filter(user=user).count() == 2

    response = auth_client.get(reverse("notifications:notification_list"))
    assert response.status_code == 200
    assert b"en retard de 10 j" in response.content
    assert "🔔" in response.content.decode()

    n = notifications.first()
    response = auth_client.get(reverse("notifications:notification_open", args=[n.id]))
    assert response.status_code == 302
    n.refresh_from_db()
    assert n.is_read is True
