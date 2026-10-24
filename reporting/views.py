"""Exports de factures (CSV / Excel) basés sur les filtres de la liste."""

import csv
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseBadRequest
from django.utils import timezone

from accounts.utils import get_user_company
from invoices.views_crud import filtered_invoices

COLUMNS = [
    ("invoice_number", "N° facture"),
    ("external_reference", "Référence externe"),
    ("client", "Client"),
    ("client_identifier", "SIRET client"),
    ("platform", "Plateforme"),
    ("platform_id", "ID plateforme"),
    ("amount_ht", "Montant HT"),
    ("vat_amount", "TVA"),
    ("amount_ttc", "Montant TTC"),
    ("currency", "Devise"),
    ("issue_date", "Date émission"),
    ("due_date", "Échéance"),
    ("payment_date", "Date paiement"),
    ("status", "Statut"),
    ("days_overdue", "Jours de retard"),
    ("is_compliant", "Conforme"),
    ("payment_method", "Mode de paiement"),
]


def invoice_row(invoice, today):
    """Ligne d'export : valeurs natives (Decimal/date) pour Excel, str pour CSV."""
    if invoice.status != "paid" and invoice.due_date < today:
        days_overdue = (today - invoice.due_date).days
    else:
        days_overdue = 0
    return [
        invoice.invoice_number,
        invoice.external_reference,
        invoice.client.name,
        invoice.client.identifier,
        invoice.get_platform_display(),
        invoice.platform_id,
        invoice.amount_ht,
        invoice.vat_amount,
        invoice.amount_ttc,
        invoice.currency.code,
        invoice.issue_date,
        invoice.due_date,
        invoice.payment_date,
        str(invoice.get_status_display()),
        days_overdue,
        "Oui" if invoice.is_compliant else "Non",
        invoice.payment_method,
    ]


def _filename(ext):
    return f"recovr_factures_{timezone.localdate():%Y%m%d}.{ext}"


def export_csv(invoices, today):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename("csv")}"'
    response.write("\ufeff")  # BOM pour Excel
    writer = csv.writer(response, delimiter=";")
    writer.writerow([label for _, label in COLUMNS])
    for invoice in invoices:
        writer.writerow(["" if v is None else v for v in invoice_row(invoice, today)])
    return response


def export_xlsx(invoices, today):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Factures"
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for invoice in invoices:
        ws.append(invoice_row(invoice, today))
    date_cols = {i + 1 for i, (key, _) in enumerate(COLUMNS) if key.endswith("date")}
    money_cols = {
        i + 1
        for i, (key, _) in enumerate(COLUMNS)
        if key in ("amount_ht", "vat_amount", "amount_ttc")
    }
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in date_cols:
                cell.number_format = "DD/MM/YYYY"
            elif cell.column in money_cols:
                cell.number_format = "#,##0.00"
    for i, (_, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(label) + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename("xlsx")}"'
    return response


@login_required
def export_invoices(request):
    company = get_user_company(request.user)
    if company is None:
        return HttpResponseBadRequest("Aucune entreprise associée à cet utilisateur.")
    fmt = request.GET.get("format", "csv").lower()
    invoices = filtered_invoices(company, request.GET)
    today = timezone.localdate()
    if fmt == "xlsx":
        return export_xlsx(invoices, today)
    if fmt == "csv":
        return export_csv(invoices, today)
    return HttpResponseBadRequest("Format inconnu (csv ou xlsx).")
